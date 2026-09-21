from __future__ import annotations

import hashlib
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from text2sql_agent.config import Settings


@dataclass(frozen=True)
class RetrievedContext:
    text: str
    sources: list[str]


class KnowledgeStore:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.embeddings = OllamaEmbeddings(
            model=settings.embedding_model,
            base_url=settings.ollama_base_url,
            client_kwargs={"timeout": settings.ollama_timeout_seconds},
        )
        settings.chroma_directory.mkdir(parents=True, exist_ok=True)
        self.vector_store = Chroma(
            collection_name=settings.chroma_collection,
            embedding_function=self.embeddings,
            persist_directory=str(settings.chroma_directory),
        )

    def retrieve(self, query: str, *, top_k: int = 6) -> RetrievedContext:
        documents = self.vector_store.similarity_search(query, k=top_k)
        sources: list[str] = []
        blocks: list[str] = []
        for index, document in enumerate(documents, start=1):
            source = str(document.metadata.get("source", "unknown"))
            location = document.metadata.get("location")
            label = f"{source} ({location})" if location else source
            if label not in sources:
                sources.append(label)
            blocks.append(
                f'<context id="{index}" source="{label}">\n{document.page_content}\n</context>'
            )
        return RetrievedContext(text="\n\n".join(blocks), sources=sources)

    def count(self) -> int:
        return int(self.vector_store._collection.count())


def load_knowledge_documents(
    directory: Path, *, patterns: tuple[str, ...] = ("*.pdf", "*.docx")
) -> list[Document]:
    documents: list[Document] = []
    paths = sorted({path for pattern in patterns for path in directory.glob(pattern)})
    for path in paths:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            documents.extend(_load_pdf(path))
        elif suffix == ".docx":
            documents.extend(_load_docx(path))
    return documents


def split_documents(
    documents: list[Document], *, chunk_size: int = 1400, chunk_overlap: int = 150
) -> tuple[list[Document], list[str]]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n## ", "\n### ", "\n\n", "\n", ". ", " "],
    )
    chunks = splitter.split_documents(documents)
    identifiers: list[str] = []
    for index, chunk in enumerate(chunks):
        identity = "|".join(
            [
                str(chunk.metadata.get("source", "unknown")),
                str(chunk.metadata.get("location", "")),
                str(index),
                chunk.page_content,
            ]
        )
        identifiers.append(hashlib.sha256(identity.encode("utf-8")).hexdigest())
    return chunks, identifiers


def index_knowledge(
    settings: Settings,
    *,
    reset: bool = False,
    progress: Callable[[int, int], None] | None = None,
    patterns: tuple[str, ...] = ("*.pdf", "*.docx"),
) -> int:
    store = KnowledgeStore(settings)
    if reset:
        store.vector_store.reset_collection()
    documents = load_knowledge_documents(settings.knowledge_directory, patterns=patterns)
    chunks, identifiers = split_documents(documents)
    total = len(chunks)
    existing = set(store.vector_store.get(include=[]).get("ids", []))
    pending = [
        (chunk, identifier)
        for chunk, identifier in zip(chunks, identifiers, strict=True)
        if identifier not in existing
    ]
    batch_size = settings.embedding_batch_size
    completed = total - len(pending)
    if progress:
        progress(completed, total)
    for start in range(0, len(pending), batch_size):
        batch = pending[start : start + batch_size]
        batch_chunks = [item[0] for item in batch]
        batch_ids = [item[1] for item in batch]
        store.vector_store.add_documents(batch_chunks, ids=batch_ids)
        completed += len(batch)
        if progress:
            progress(completed, total)
    return len(chunks)


def _load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(path)
    documents: list[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={"source": path.name, "location": f"page {page_number}"},
                )
            )
    return documents


def _load_docx(path: Path) -> list[Document]:
    document = DocxDocument(path)
    sections: list[Document] = []
    current_heading = "document start"
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            sections.append(
                Document(
                    page_content="\n".join(buffer).strip(),
                    metadata={"source": path.name, "location": current_heading},
                )
            )
            buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            flush()
            current_heading = text
        buffer.append(text)
    flush()
    return sections
